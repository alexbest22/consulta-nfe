import logging
import re
import time
from datetime import date, timedelta
from pathlib import Path
from typing import List, Optional, Tuple

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image
from pytesseract import pytesseract

pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
import pytesseract
from selenium.common.exceptions import TimeoutException
from selenium.webdriver import Chrome
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait

from config import ELEMENT_WAIT_SECONDS

try:
    import pyautogui

    pyautogui.FAILSAFE = False
except ImportError:  # pragma: no cover - ambiente sem pyautogui
    pyautogui = None  # type: ignore[assignment]


def configurar_logs() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="[%(asctime)s] %(levelname)s - %(message)s",
        datefmt="%H:%M:%S",
    )


def pausar(segundos: float) -> None:
    if segundos > 0:
        logging.debug("Pausa de %.2f segundos.", segundos)
        time.sleep(segundos)


def carregar_credenciais(
    xlsx_path: Path,
    sheet_name: Optional[str],
    row_number: int,
    col_user: int,
    col_password: int,
) -> Tuple[str, str]:
    logging.info("STEP 1 - Lendo credenciais em %s", xlsx_path)
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    usuario = worksheet.cell(row=row_number, column=col_user).value
    senha = worksheet.cell(row=row_number, column=col_password).value
    if not usuario or not senha:
        raise ValueError(
            "Nao foi possivel encontrar usuario e senha na planilha. "
            "Verifique os parametros --linha, --coluna-usuario e --coluna-senha."
        )
    return str(usuario).strip(), str(senha).strip()


def carregar_lista_credenciais(
    xlsx_path: Path,
    sheet_name: Optional[str],
    row_start: int,
    col_user: int,
    col_password: int,
    col_codigo: int,
) -> List[Tuple[str, str, str]]:
    logging.info("STEP 1 - Lendo todas as credenciais em %s", xlsx_path)
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    credenciais: List[Tuple[str, str, str]] = []
    linha = row_start
    while True:
        codigo = worksheet.cell(row=linha, column=col_codigo).value
        usuario = worksheet.cell(row=linha, column=col_user).value
        senha = worksheet.cell(row=linha, column=col_password).value
        if codigo and usuario and senha:
            credenciais.append(
                (
                    str(codigo).strip(),
                    str(usuario).strip(),
                    str(senha).strip(),
                )
            )
        else:
            break
        linha += 1
    return credenciais


def periodo_mes_anterior(base: Optional[date] = None) -> Tuple[str, str]:
    hoje = base or date.today()
    primeiro_deste_mes = hoje.replace(day=1)
    ultimo_mes = primeiro_deste_mes - timedelta(days=1)
    primeiro = ultimo_mes.replace(day=1)
    ultimo = ultimo_mes
    inicio = primeiro.strftime("%d/%m/%Y")
    fim = ultimo.strftime("%d/%m/%Y")
    logging.info("STEP 2 - Periodo selecionado: %s ate %s (mes anterior).", inicio, fim)
    return inicio, fim


MESES_PT = [
    "janeiro",
    "fevereiro",
    "marco",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def periodo_mes_anterior_info(base: Optional[date] = None) -> Tuple[str, str, int, str]:
    hoje = base or date.today()
    primeiro_deste_mes = hoje.replace(day=1)
    ultimo_mes = primeiro_deste_mes - timedelta(days=1)
    primeiro = ultimo_mes.replace(day=1)
    inicio = primeiro.strftime("%d/%m/%Y")
    fim = ultimo_mes.strftime("%d/%m/%Y")
    ano = ultimo_mes.year
    mes_extenso = MESES_PT[ultimo_mes.month - 1]
    return inicio, fim, ano, mes_extenso


def pressionar_enter(headless: bool, delay: float = 0.0, contexto: str = "") -> None:
    if headless:
        logging.info("STEP - Modo headless ativo; pulando envio de ENTER (%s).", contexto)
        return
    if pyautogui is None:
        logging.info("STEP - pyautogui nao instalado; pulando envio de ENTER (%s).", contexto)
        return

    if delay > 0:
        time.sleep(delay)
    try:
        pyautogui.press("enter")
        logging.debug("STEP - ENTER enviado via pyautogui (%s).", contexto)
    except Exception as exc:
        logging.warning("STEP - Falha ao enviar ENTER (%s): %s", contexto, exc)


def detectar_captcha(driver: Chrome) -> bool:
    seletor_xpath = (
        "//*[contains(translate(@id, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'captcha')"
        " or contains(translate(@class, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'captcha')"
        " or contains(translate(@src, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'captcha')]"
    )
    captcha_elements = driver.find_elements(By.XPATH, seletor_xpath)
    return len(captcha_elements) > 0


def resolver_captcha_ocr(driver: Chrome) -> bool:
    from selectors_module import DEFAULT_SELECTORS

    logging.info("STEP - Tentando resolver CAPTCHA por OCR.")
    try:
        imagem_elemento = driver.find_element(*DEFAULT_SELECTORS.captcha_image[0])
        campo_entrada = driver.find_element(*DEFAULT_SELECTORS.captcha_input[0])
    except Exception:
        logging.info("STEP - Elementos de CAPTCHA nao encontrados para OCR.")
        return False

    imagem_b64 = imagem_elemento.screenshot_as_png
    texto = extrair_texto_captcha_from_bytes(imagem_b64)
    if texto:
        logging.info("STEP - OCR retornou '%s'.", texto)
        preencher_campo(driver, campo_entrada, texto)
        return True

    logging.warning("STEP - OCR nao conseguiu decodificar o CAPTCHA.")
    return False


def lidar_com_captcha(driver: Chrome, destino: Path, aguardar_usuario: bool) -> None:
    logging.info("STEP - Verificando presenca de CAPTCHA.")
    if not detectar_captcha(driver):
        return

    logging.warning(
        "STEP - CAPTCHA detectado. Informe manualmente caso o OCR nao esteja ativado."
    )
    if aguardar_usuario:
        input("Digite o CAPTCHA no navegador e pressione Enter para continuar...")


def extrair_texto_captcha_from_bytes(image_bytes: bytes) -> str:
    np_array = np.frombuffer(image_bytes, np.uint8)
    imagem = cv2.imdecode(np_array, cv2.IMREAD_GRAYSCALE)
    if imagem is None:
        return ""

    blur = cv2.medianBlur(imagem, 3)
    _, thresh = cv2.threshold(blur, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    kernel = np.ones((2, 2), np.uint8)
    close = cv2.morphologyEx(thresh, cv2.MORPH_CLOSE, kernel, iterations=2)
    dilate = cv2.dilate(close, kernel, iterations=1)
    erode = cv2.erode(dilate, kernel, iterations=1)
    ampliada = cv2.resize(erode, None, fx=2, fy=2, interpolation=cv2.INTER_LINEAR)

    pil_image = Image.fromarray(ampliada)
    texto = pytesseract.image_to_string(
        pil_image,
        config="--psm 8 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789",
    )
    formatado = "".join(ch for ch in texto.upper() if ch.isalnum())
    return formatado[:4]


def aguardar_download_conclusao(
    diretorio: Path, timeout: int = 180, estabilidade: int = 2
) -> Path:
    logging.info("STEP - Aguardando conclusao do download na pasta %s.", diretorio)
    fim = time.time() + timeout
    diretorio.mkdir(parents=True, exist_ok=True)
    ultimo_arquivo: Optional[Path] = None
    ultimo_tamanho: Optional[int] = None
    tempo_estavel = time.time()

    while time.time() < fim:
        cr = list(diretorio.glob("*.crdownload"))
        arquivos = [p for p in diretorio.iterdir() if p.is_file()]

        if arquivos:
            candidato = max(arquivos, key=lambda p: p.stat().st_mtime)
            tamanho = candidato.stat().st_size

            if cr:
                ultimo_arquivo = None
                ultimo_tamanho = None
                tempo_estavel = time.time()
            else:
                if candidato != ultimo_arquivo:
                    ultimo_arquivo = candidato
                    ultimo_tamanho = tamanho
                    tempo_estavel = time.time()
                elif tamanho != ultimo_tamanho:
                    ultimo_tamanho = tamanho
                    tempo_estavel = time.time()
                elif time.time() - tempo_estavel >= estabilidade:
                    logging.info("STEP - Download concluido: %s", candidato.name)
                    return candidato

        time.sleep(1)

    raise TimeoutError("Download nao concluido dentro do tempo limite.")


def normalizar_nome_pasta(texto: str, fallback: str) -> str:
    if not texto:
        texto = fallback
    texto = texto.strip()
    texto = re.sub(r"[^0-9A-Za-zÀ-ÿ\s_-]", "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto[:80] if texto else fallback


def ajustar_extensao_planilha(arquivo: Path, extensao_desejada: str = ".xls") -> Path:
    if arquivo.suffix.lower() == ".tmp":
        novo_arquivo = arquivo.with_suffix(extensao_desejada)
        logging.info(
            "STEP - Renomeando arquivo temporario %s para %s",
            arquivo.name,
            novo_arquivo.name,
        )
        arquivo.rename(novo_arquivo)
        return novo_arquivo
    return arquivo


def aguardar_pagina_carregada(driver: Chrome, timeout: int) -> None:
    WebDriverWait(driver, timeout).until(
        lambda d: d.execute_script("return document.readyState") == "complete"
    )
    logging.info(
        "STEP 4 - Estado do documento apos carregamento: %s",
        driver.execute_script("return document.readyState"),
    )


def preencher_campo(driver: Chrome, elemento, valor: str) -> None:
    try:
        elemento.click()
        elemento.clear()
        elemento.send_keys(valor)
    except Exception as exc:
        logging.warning(
            "Falha ao preencher campo via send_keys (%s). Aplicando fallback JavaScript.",
            exc,
        )
        driver.execute_script(
            (
                "arguments[0].value = arguments[1];"
                "arguments[0].dispatchEvent(new Event('input', {bubbles: true}));"
                "arguments[0].dispatchEvent(new Event('change', {bubbles: true}));"
            ),
            elemento,
            valor,
        )
    else:
        driver.execute_script(
            (
                "arguments[0].value = arguments[1];"
                "arguments[0].dispatchEvent(new Event('input', {bubbles: true}));"
                "arguments[0].dispatchEvent(new Event('change', {bubbles: true}));"
            ),
            elemento,
            valor,
        )

    def _tem_valor(_: object) -> bool:
        try:
            return (elemento.get_attribute("value") or "").strip() == valor
        except Exception:
            return False

    WebDriverWait(driver, ELEMENT_WAIT_SECONDS).until(_tem_valor)
